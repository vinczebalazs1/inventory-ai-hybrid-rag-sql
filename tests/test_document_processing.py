import json
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from pydantic import ValidationError

from app.document_processing import processing_service as service
from app.document_processing.config import Settings
from app.document_processing.docling_parser import parse_document
from app.document_processing.extraction_service import chunks, extract_fields
from app.document_processing.field_mapper import map_fields
from app.document_processing.models import (
    ExtractedField, ExtractionResult, MappedField, MappingResult, ParsedDocument,
)
from app.document_processing.validation_service import normalize, validate_fields
from app.errors import AppError
from app.services import llm_service


def parsed(text="Location: Budapest", pages=None):
    return ParsedDocument(source_name="test.pdf", source_extension=".pdf",
                          markdown=text, extraction_text=text, pages=pages or {})


def mock_llm(monkeypatch, *responses):
    mock = Mock(side_effect=[json.dumps(r) if isinstance(r, dict) else r for r in responses])
    monkeypatch.setattr(llm_service, "ask_llm", mock)
    return mock


def payload(**kwargs):
    return {"fields": [{"label": "Location", "value": "Budapest",
                         "evidence": "Location: Budapest", **kwargs}]}


def test_generic_extraction_and_provenance(monkeypatch):
    mock_llm(monkeypatch, payload(page=99))
    result = extract_fields(parsed(pages={2: "Location: Budapest"}), Settings())
    assert result.fields[0].page == 2
    assert result.fields[0].evidence == "Location: Budapest"


def test_one_repair(monkeypatch):
    mock = mock_llm(monkeypatch, "invalid", payload())
    assert extract_fields(parsed(), Settings()).fields[0].value == "Budapest"
    assert mock.call_count == 2
    assert "invalid" in mock.call_args.args[0]
    assert "Validation errors" in mock.call_args.args[0]


def test_repair_is_bounded(monkeypatch):
    mock = mock_llm(monkeypatch, "invalid", "still invalid")
    with pytest.raises(ValidationError):
        extract_fields(parsed(), Settings())
    assert mock.call_count == 2


@pytest.mark.parametrize("kwargs", [{"label": ""}, {"label": " "}, {"page": 0}, {"value": 42}, {"page": True}])
def test_field_validation(kwargs):
    with pytest.raises(ValidationError):
        ExtractedField(**{"label": "A", "value": "B", **kwargs})


def test_false_evidence_and_page_removed(monkeypatch):
    mock_llm(monkeypatch, payload(evidence="invented", page=1))
    field = extract_fields(parsed(), Settings()).fields[0]
    assert field.evidence is None and field.page is None


def test_docx_does_not_get_invented_page(monkeypatch):
    mock_llm(monkeypatch, payload(page=1))
    assert extract_fields(parsed(), Settings()).fields[0].page is None


def test_merge_duplicates_keeps_different_values(monkeypatch):
    content = "Location: Budapest\n\nLocation: London\n\n" * 4
    response = {"fields": [payload()["fields"][0], {"label": "Location", "value": "London", "evidence": "Location: London"}]}
    count = len(list(chunks(content, 80, 30)))
    mock_llm(monkeypatch, *([response] * count))
    fields = extract_fields(parsed(content), Settings(chunk_size=80, chunk_overlap=30)).fields
    assert {f.value for f in fields} == {"Budapest", "London"}
    assert len(fields) == 2


def test_chunk_size_coverage_and_overlap():
    text = "paragraph text\n\n" * 100
    result = list(chunks(text, 80, 15))
    assert all(len(chunk) <= 80 for chunk in result)
    assert result[0].startswith("paragraph")
    assert result[-1].endswith("\n\n")
    with pytest.raises(ValueError):
        list(chunks(text, 10, 10))


def test_duplicate_across_pages_keeps_first_page(monkeypatch):
    mock_llm(monkeypatch, payload(), payload())
    result = extract_fields(parsed(pages={1: "Location: Budapest", 2: "Location: Budapest"}), Settings())
    assert len(result.fields) == 1 and result.fields[0].page == 1


def test_schema_load_independent_of_cwd(monkeypatch, tmp_path):
    monkeypatch.chdir(tmp_path)
    assert "inventory_items" in service.schema_service.load_schema()


def test_mapping_known_unknown_and_identifiers(monkeypatch):
    labels = ["Location", "Inventory No.", "Service Tag", "Serial Number", "Warranty", "Manufacturer"]
    extraction = ExtractionResult(source_name="a", fields=[ExtractedField(label=x, value="x") for x in labels])
    mock = mock_llm(monkeypatch, {"assignments": [
        {"index": i, "target_field": target} for i, target in enumerate(["location", "inventory_number", None, None, None, None])
    ]})
    result = map_fields(extraction, [{"column": "location", "type": "text"},
                                    {"column": "inventory_number", "type": "text"}], {}, Settings())
    assert [f.target_field for f in result.mapped_fields] == ["location", "inventory_number"]
    assert [f.label for f in result.unmapped_fields] == labels[2:]
    assert "Service Tag and Serial Number are not internal inventory numbers" in mock.call_args.args[0]


@pytest.mark.parametrize("assignments", [
    [{"index": 0, "target_field": "invented"}], [],
    [{"index": 1, "target_field": None}],
    [{"index": 0, "target_field": None}, {"index": 0, "target_field": None}],
])
def test_invalid_mapping_rejected(monkeypatch, assignments):
    mock_llm(monkeypatch, {"assignments": assignments})
    with pytest.raises(ValueError):
        map_fields(ExtractionResult(source_name="a", fields=[ExtractedField(label="x", value="y")]),
                   [{"column": "location", "type": "text"}], {}, Settings())


def test_new_schema_column_and_batches(monkeypatch):
    mock = mock_llm(monkeypatch, *([{"assignments": [{"index": 0, "target_field": "new_field"}]}] * 2))
    result = map_fields(ExtractionResult(source_name="a", fields=[ExtractedField(label="x", value=str(i)) for i in range(2)]),
                        [{"column": "new_field", "type": "text"}], {"new_field": "New meaning"},
                        Settings(mapping_batch_size=1))
    assert len(result.mapped_fields) == 2 and mock.call_count == 2
    assert "New meaning" in mock.call_args.args[0]


@pytest.mark.parametrize("value,dtype,field,expected", [
    ("1 234", "integer", "quantity", 1234),
    ("1.234,56", "numeric", "purchase_value", "1234.56"),
    ("1234.56", "decimal", "depreciation", "1234.56"),
    ("2024.02.29.", "date", "activation_date", "2024-02-29"),
    ("29.02.2024", "date", "activation_date", "2024-02-29"),
    (" Ft ", "text", "currency", "HUF"),
    (" x ", "text", "name", "x"),
])
def test_normalization(value, dtype, field, expected):
    assert normalize(value, dtype, field) == expected


@pytest.mark.parametrize("value,dtype,field", [
    ("1.2", "integer", "quantity"), ("NaN", "numeric", "purchase_value"),
    ("1,234", "numeric", "current_value"), ("2023-02-29", "date", "activation_date"),
    ("01/02/2024", "date", "activation_date"), ("$", "text", "currency"), (" ", "text", "name"),
])
def test_invalid_value_preserves_raw(value, dtype, field):
    mapping = MappingResult(mapped_fields=[MappedField(source_label="x", raw_value=value, target_field=field)])
    issues = validate_fields(mapping, [{"column": field, "type": dtype}])
    assert len(issues) == 1
    assert mapping.mapped_fields[0].raw_value == value
    assert mapping.mapped_fields[0].normalized_value is None


def test_orchestration(monkeypatch):
    monkeypatch.setattr(service, "parse_document", lambda *a, **kw: parsed())
    mock_llm(monkeypatch, payload(), {"assignments": [{"index": 0, "target_field": "location"}]})
    result = service.process_document(b"data", "a.pdf")
    assert result.mapping.mapped_fields[0].normalized_value == "Budapest"
    assert not result.validation_issues
    assert json.loads(result.model_dump_json())["extraction"]["source_name"] == "test.pdf"


def test_stage_error_is_friendly(monkeypatch):
    monkeypatch.setattr(service, "parse_document", lambda *a, **kw: parsed())
    mock_llm(monkeypatch, "invalid", "invalid")
    with pytest.raises(AppError, match="LLM mezőkinyerés"):
        service.process_document(b"x", "a.pdf")


def test_unsupported_extension():
    with pytest.raises(AppError, match="PDF vagy DOCX"):
        parse_document(b"x", "../../bad.exe")


def test_converter_adapter():
    from docling.datamodel.base_models import ConversionStatus
    document = Mock()
    document.pages = {1: object()}
    document.export_to_markdown.return_value = "Location: Budapest"
    document.export_to_dict.return_value = {"name": "demo"}
    converter = Mock()
    converter.convert.return_value = SimpleNamespace(status=ConversionStatus.SUCCESS, document=document)
    result = parse_document(b"pdf", "../../demo.pdf", converter=converter)
    assert result.source_name == "demo.pdf"
    assert result.pages == {1: "Location: Budapest"}
    assert "[PAGE 1]" in result.extraction_text
    assert converter.convert.call_args.args[0].stream.getvalue() == b"pdf"


def test_real_docx_smoke(monkeypatch):
    from docx import Document
    document = Document()
    document.add_paragraph("Inventory No.: INV-123")
    document.add_paragraph("Location: Budapest")
    buffer = BytesIO()
    document.save(buffer)
    result = parse_document(buffer.getvalue(), "demo.docx")
    assert "INV-123" in result.markdown
    assert "Budapest" in result.extraction_text
    assert result.raw_document and result.pages == {}
    mock_llm(monkeypatch, payload(), {"assignments": [{"index": 0, "target_field": "location"}]})
    processed = service.process_document(buffer.getvalue(), "demo.docx")
    assert processed.mapping.mapped_fields[0].normalized_value == "Budapest"
    assert processed.mapping.mapped_fields[0].evidence == "Location: Budapest"
    assert processed.mapping.mapped_fields[0].page is None


def test_streamlit_startup():
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_file("app/web/simple_server.py").run(timeout=30)
    assert not app.exception
    assert any(button.label == "Analyze document" for button in app.button)


def test_streamlit_preview_tables(monkeypatch):
    from streamlit.testing.v1 import AppTest
    monkeypatch.setattr(service, "parse_document", lambda *a, **kw: parsed())
    mock_llm(monkeypatch, {"fields": [
        {"label": "Location", "value": "Budapest", "evidence": "Location: Budapest"},
        {"label": "Quantity", "value": "bad"}, {"label": "Warranty", "value": "36 months"},
    ]}, {"assignments": [
        {"index": 0, "target_field": "location"}, {"index": 1, "target_field": "quantity"},
        {"index": 2, "target_field": None},
    ]})
    result = service.process_document(b"x", "demo.pdf")
    app = AppTest.from_file("app/web/simple_server.py")
    app.session_state["document_preview"] = result
    app.run(timeout=30)
    assert not app.exception
    assert len(app.dataframe) == 3
    assert "evidence" in app.dataframe[0].value.columns
    assert len(app.json) == 2
